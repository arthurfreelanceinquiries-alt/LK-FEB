import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os
import subprocess

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{kwargs.get("top", "single")}" w:sz="{kwargs.get("top_sz", "4")}" w:space="0" w:color="{kwargs.get("top_color", "000000")}"/>\n'
        f'  <w:left w:val="{kwargs.get("left", "single")}" w:sz="{kwargs.get("left_sz", "4")}" w:space="0" w:color="{kwargs.get("left_color", "000000")}"/>\n'
        f'  <w:bottom w:val="{kwargs.get("bottom", "single")}" w:sz="{kwargs.get("bottom_sz", "4")}" w:space="0" w:color="{kwargs.get("bottom_color", "000000")}"/>\n'
        f'  <w:right w:val="{kwargs.get("right", "single")}" w:sz="{kwargs.get("right_sz", "4")}" w:space="0" w:color="{kwargs.get("right_color", "000000")}"/>\n'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def create_excel_invoice():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice BEM FEB"
    
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_view.showGridLines = True
    
    font_family = "Arial"
    font_kop_univ = Font(name="Times New Roman", size=13, bold=True)
    font_kop_bem = Font(name="Arial", size=10, bold=True)
    font_kop_addr = Font(name="Times New Roman", size=8.5)
    font_title = Font(name=font_family, size=16, bold=True)
    font_header_bold = Font(name=font_family, size=10.5, bold=True)
    font_regular = Font(name=font_family, size=10)
    font_bold = Font(name=font_family, size=10, bold=True)
    font_table_head = Font(name=font_family, size=10, bold=True)
    
    peach_fill = PatternFill(start_color="F5D5D3", end_color="F5D5D3", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )
    
    thick_bottom_border = Border(
        bottom=Side(style='medium', color='000000')
    )

    # Insert Left Logo (UKRIDA)
    if os.path.exists("ukrida_logo_tight.png"):
        img_l = OpenpyxlImage("ukrida_logo_tight.png")
        img_l.width = 72
        img_l.height = 90
        ws.add_image(img_l, "A1")
        
    # Insert Right Logo (BEM FEB)
    if os.path.exists("BEM_FEB_LOGO_tight.png"):
        img_r = OpenpyxlImage("BEM_FEB_LOGO_tight.png")
        img_r.width = 77
        img_r.height = 90
        ws.add_image(img_r, "E1")
        
    # Kop Surat Center Text (Rows 1-3)
    ws.merge_cells("B1:D1")
    ws["B1"] = "UNIVERSITAS KRISTEN KRIDA WACANA"
    ws["B1"].font = font_kop_univ
    ws["B1"].alignment = Alignment(horizontal="center", vertical="center")
    
    ws.merge_cells("B2:D2")
    ws["B2"] = "BADAN EKSEKUTIF MAHASISWA FAKULTAS EKONOMI & BISNIS"
    ws["B2"].font = font_kop_bem
    ws["B2"].alignment = Alignment(horizontal="center", vertical="center")
    
    ws.merge_cells("B3:D3")
    ws["B3"] = "Jl. Tanjung Duren Raya No.4 Telp. 5666952-189 Fax 5666956, Jakarta Barat 11470"
    ws["B3"].font = font_kop_addr
    ws["B3"].alignment = Alignment(horizontal="center", vertical="center")
    
    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 16
    ws.row_dimensions[4].height = 8
    
    # Divider line
    for col in range(1, 6):
        ws.cell(row=4, column=col).border = thick_bottom_border
        
    # Title
    ws["A6"] = "INVOICE"
    ws["A6"].font = font_title
    
    # Metadata
    ws["D6"] = "No. Invoice: INV/BEM-FEB/OM/IX/2026"
    ws["D6"].font = font_regular
    ws["D7"] = "Tanggal: 08 September 2026"
    ws["D7"].font = font_regular
    
    # Recipient
    ws["A8"] = "DITUJUKAN KEPADA:"
    ws["A8"].font = font_header_bold
    
    ws["A9"] = "PT. Inti Citra Agung (Optik Melawai)"
    ws["A9"].font = font_regular
    ws["A10"] = "Jl. KS Tubun Raya No. 64"
    ws["A10"].font = font_regular
    ws["A11"] = "RT 001 RW 001"
    ws["A11"].font = font_regular
    ws["A12"] = "Slipi Palmerah"
    ws["A12"].font = font_regular
    ws["A13"] = "Jakarta Barat 11410"
    ws["A13"].font = font_regular
    
    # Event Note
    ws["A15"] = "Kegiatan: BEM FEB X Optik Melawai 2026 (RPK Periode 2026/2027)"
    ws["A15"].font = Font(name=font_family, size=9.5, italic=True)
    
    # Table Header at Row 17
    headers = ["Keterangan", "Harga", "Hari/Tanggal", "Jumlah", "Total"]
    ws.row_dimensions[17].height = 24
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=17, column=col_idx, value=h)
        cell.font = font_table_head
        cell.fill = peach_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    # Table Content (Single revised item: Sumbangan 3 Juta)
    items = [
        ("Sumbangan", 3000000, "10-11 September 2026", 1, 3000000),
    ]
    
    start_row = 18
    for idx, (ket, harga, tgl, jml, tot) in enumerate(items):
        r = start_row + idx
        ws.row_dimensions[r].height = 24
        c_ket = ws.cell(row=r, column=1, value=ket)
        c_ket.font = font_regular
        c_ket.alignment = Alignment(horizontal="left", vertical="center")
        
        c_harga = ws.cell(row=r, column=2, value=harga)
        c_harga.font = font_regular
        c_harga.alignment = Alignment(horizontal="right", vertical="center")
        c_harga.number_format = '"Rp"#,##0.00'
        
        c_tgl = ws.cell(row=r, column=3, value=tgl)
        c_tgl.font = font_regular
        c_tgl.alignment = Alignment(horizontal="center", vertical="center")
        
        c_jml = ws.cell(row=r, column=4, value=jml)
        c_jml.font = font_regular
        c_jml.alignment = Alignment(horizontal="center", vertical="center")
        
        c_tot = ws.cell(row=r, column=5, value=f"=B{r}*D{r}")
        c_tot.font = font_regular
        c_tot.alignment = Alignment(horizontal="right", vertical="center")
        c_tot.number_format = '"Rp"#,##0.00'
        
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = thin_border

    # Total Row
    total_row = start_row + len(items)
    ws.row_dimensions[total_row].height = 24
    ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=4)
    cell_tot_lbl = ws.cell(row=total_row, column=1, value="Total")
    cell_tot_lbl.font = font_bold
    cell_tot_lbl.alignment = Alignment(horizontal="center", vertical="center")
    
    cell_tot_val = ws.cell(row=total_row, column=5, value=f"=SUM(E{start_row}:E{total_row-1})")
    cell_tot_val.font = font_bold
    cell_tot_val.alignment = Alignment(horizontal="right", vertical="center")
    cell_tot_val.number_format = '"Rp"#,##0.00'
    
    for c in range(1, 6):
        ws.cell(row=total_row, column=c).border = thin_border
        
    # Terbilang
    ws.cell(row=total_row+1, column=1, value="Terbilang: Tiga Juta Rupiah").font = Font(name=font_family, size=9.5, italic=True)
    
    # Payment Info
    p_row = total_row + 3
    ws.cell(row=p_row, column=1, value="Pembayaran:").font = font_bold
    ws.cell(row=p_row+1, column=1, value="Nama Rekening:").font = font_regular
    ws.cell(row=p_row+1, column=2, value="JEMIMA PUTRI SUTANTA").font = font_bold
    
    ws.cell(row=p_row+2, column=1, value="Nomor Rekening:").font = font_regular
    ws.cell(row=p_row+2, column=2, value="8400364932").font = font_bold
    
    ws.cell(row=p_row+3, column=1, value="Bank:").font = font_regular
    ws.cell(row=p_row+3, column=2, value="BCA").font = font_bold
    
    # Signature block
    s_row = p_row + 5
    ws.cell(row=s_row, column=1, value="Hormat kami,").font = font_regular
    ws.cell(row=s_row+1, column=1, value="Ketua BEM FEB UKRIDA").font = font_bold
    
    ws.cell(row=s_row, column=4, value="Mengetahui,").font = font_regular
    ws.cell(row=s_row+1, column=4, value="Bendahara BEM FEB UKRIDA").font = font_bold
    
    ws.cell(row=s_row+5, column=1, value="(Arthur Reezan)").font = font_bold
    ws.cell(row=s_row+6, column=1, value="NIM: 312023002").font = font_regular
    
    ws.cell(row=s_row+5, column=4, value="(Jemima Putri Sutanta)").font = font_bold
    ws.cell(row=s_row+6, column=4, value="NIM: 322024034").font = font_regular

    # Column widths
    ws.column_dimensions['A'].width = 38
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 20
    
    wb.save("Invoice_BEM_FEB_Optik_Melawai.xlsx")
    print("Saved Invoice_BEM_FEB_Optik_Melawai.xlsx with dual logos.")

def create_word_invoice():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        section.page_width = Inches(8.27)  # A4
        section.page_height = Inches(11.69)
        
    # Header / Kop Surat table with 3 columns (Left Logo, Center Text, Right Logo)
    header_table = doc.add_table(rows=1, cols=3)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_table.autofit = False
    header_table.columns[0].width = Inches(1.0)
    header_table.columns[1].width = Inches(4.87)
    header_table.columns[2].width = Inches(1.0)
    
    # Left Logo Cell (UKRIDA)
    cell_l = header_table.cell(0, 0)
    p_l = cell_l.paragraphs[0]
    p_l.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if os.path.exists("ukrida_logo_tight.png"):
        p_l.add_run().add_picture("ukrida_logo_tight.png", width=Inches(0.95))
        
    # Center Text Cell
    cell_c = header_table.cell(0, 1)
    p_c = cell_c.paragraphs[0]
    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c.paragraph_format.line_spacing = 1.15
    
    r1 = p_c.add_run("UNIVERSITAS KRISTEN KRIDA WACANA\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(13)
    r1.bold = True
    
    r2 = p_c.add_run("BADAN EKSEKUTIF MAHASISWA FAKULTAS EKONOMI & BISNIS\n")
    r2.font.name = "Arial"
    r2.font.size = Pt(9.5)
    r2.bold = True
    
    r3 = p_c.add_run("Jl. Tanjung Duren Raya No.4 Telp. 5666952-189 Fax 5666956, Jakarta Barat 11470")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(8)
    
    # Right Logo Cell (BEM FEB)
    cell_r = header_table.cell(0, 2)
    p_r = cell_r.paragraphs[0]
    p_r.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if os.path.exists("BEM_FEB_LOGO_tight.png"):
        p_r.add_run().add_picture("BEM_FEB_LOGO_tight.png", width=Inches(0.98))
        
    # Solid black line under kop
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_before = Pt(4)
    p_line.paragraph_format.space_after = Pt(10)
    run_line = p_line.add_run("―" * 58)
    run_line.font.name = "Arial"
    run_line.font.size = Pt(14)
    run_line.bold = True
    
    # Title & Metadata table
    meta_table = doc.add_table(rows=1, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.columns[0].width = Inches(3.5)
    meta_table.columns[1].width = Inches(3.37)
    
    # Left: INVOICE
    p_title = meta_table.cell(0, 0).paragraphs[0]
    r_title = p_title.add_run("INVOICE")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(16)
    r_title.bold = True
       # Right: Meta
    p_meta_r = meta_table.cell(0, 1).paragraphs[0]
    p_meta_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_meta_r.paragraph_format.line_spacing = 1.15
    r_meta_no = p_meta_r.add_run("No. Invoice: INV/BEM-FEB/OM/IX/2026\nTanggal: 08 September 2026\nHal: Sumbangan Kegiatan")
    r_meta_no.font.name = "Arial"
    r_meta_no.font.size = Pt(9)
    
    # Recipient
    p_to = doc.add_paragraph()
    p_to.paragraph_format.line_spacing = 1.15
    p_to.paragraph_format.space_before = Pt(8)
    p_to.paragraph_format.space_after = Pt(10)
    
    r_to_head = p_to.add_run("DITUJUKAN KEPADA:\n")
    r_to_head.bold = True
    r_to_head.font.name = "Arial"
    r_to_head.font.size = Pt(10)
    
    to_text = (
        "PT. Inti Citra Agung (Optik Melawai)\n"
        "Jl. KS Tubun Raya No. 64\n"
        "RT 001 RW 001\n"
        "Slipi Palmerah\n"
        "Jakarta Barat 11410"
    )
    r_to_body = p_to.add_run(to_text)
    r_to_body.font.name = "Arial"
    r_to_body.font.size = Pt(9.5)
    
    # Activity note
    p_act = doc.add_paragraph()
    p_act.paragraph_format.space_after = Pt(8)
    r_act = p_act.add_run("Kegiatan: BEM FEB X Optik Melawai 2026 (RPK Periode 2026/2027)")
    r_act.font.name = "Arial"
    r_act.font.size = Pt(9)
    r_act.italic = True
    
    # Table (5 columns: Keterangan, Harga, Hari/Tanggal, Jumlah, Total)
    # 3 rows: Header (0), Item (1), Total (2)
    table = doc.add_table(rows=3, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    col_widths = [Inches(2.5), Inches(1.3), Inches(1.2), Inches(0.6), Inches(1.4)]
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width
            
    # Headers
    headers = ["Keterangan", "Harga", "Hari/Tanggal", "Jumlah", "Total"]
    for idx, h in enumerate(headers):
        cell = table.cell(0, idx)
        cell.text = h
        set_cell_shading(cell, "F5D5D3")
        set_cell_border(cell, top="single", bottom="single", left="single", right="single")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.bold = True

    # Single Revised Data Row (Row 1): Sumbangan 3 Juta
    items_data = [
        ("Sumbangan", "Rp 3.000.000,00", "10-11 September 2026", "1", "Rp 3.000.000,00")
    ]
    
    for row_idx, item in enumerate(items_data, start=1):
        for col_idx, val in enumerate(item):
            cell = table.cell(row_idx, col_idx)
            cell.text = val
            set_cell_border(cell, top="single", bottom="single", left="single", right="single")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif col_idx in [1, 4]:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Arial"
                r.font.size = Pt(9)
                
    # Total Row (Row 2)
    table.cell(2, 0).merge(table.cell(2, 3))
    tot_lbl_cell = table.cell(2, 0)
    tot_lbl_cell.text = "Total"
    tot_lbl_p = tot_lbl_cell.paragraphs[0]
    tot_lbl_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in tot_lbl_p.runs:
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.bold = True
        
    tot_val_cell = table.cell(2, 1)
    tot_val_cell.text = "Rp 3.000.000,00"
    tot_val_p = tot_val_cell.paragraphs[0]
    tot_val_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in tot_val_p.runs:
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.bold = True
        
    set_cell_border(tot_lbl_cell, top="single", bottom="single", left="single", right="single")
    set_cell_border(tot_val_cell, top="single", bottom="single", left="single", right="single")
    
    # Terbilang
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(8)
    p_sp.paragraph_format.space_after = Pt(4)
    r_tb = p_sp.add_run("Terbilang: Tiga Juta Rupiah")
    r_tb.font.name = "Arial"
    r_tb.font.size = Pt(9)
    r_tb.italic = True
    
    # Payment section
    p_pay = doc.add_paragraph()
    p_pay.paragraph_format.line_spacing = 1.15
    p_pay.paragraph_format.space_before = Pt(12)
    p_pay.paragraph_format.space_after = Pt(22)
    
    r_pay_head = p_pay.add_run("Pembayaran:\n")
    r_pay_head.bold = True
    r_pay_head.font.name = "Arial"
    r_pay_head.font.size = Pt(9.5)
    
    r_pay_body = p_pay.add_run(
        "Nama Rekening: JEMIMA PUTRI SUTANTA\n"
        "Nomor Rekening: 8400364932\n"
        "Bank: BCA"
    )
    r_pay_body.font.name = "Arial"
    r_pay_body.font.size = Pt(9)
    
    # Signatures Table
    sig_table = doc.add_table(rows=3, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    sig_table.columns[0].width = Inches(3.4)
    sig_table.columns[1].width = Inches(3.4)
    
    cell_s1 = sig_table.cell(0, 0)
    cell_s1.text = "Hormat kami,\nKetua BEM FEB UKRIDA"
    for r in cell_s1.paragraphs[0].runs:
        r.font.name = "Arial"
        r.font.size = Pt(9)
        
    cell_s2 = sig_table.cell(0, 1)
    cell_s2.text = "Mengetahui,\nBendahara BEM FEB UKRIDA"
    for r in cell_s2.paragraphs[0].runs:
        r.font.name = "Arial"
        r.font.size = Pt(9)
        
    sig_table.cell(1, 0).text = "\n\n"
    sig_table.cell(1, 1).text = "\n\n"
    
    cell_s1_name = sig_table.cell(2, 0)
    cell_s1_name.text = "(Arthur Reezan)\nNIM: 312023002"
    for r in cell_s1_name.paragraphs[0].runs:
        r.font.name = "Arial"
        r.font.size = Pt(9)
        
    cell_s2_name = sig_table.cell(2, 1)
    cell_s2_name.text = "(Jemima Putri Sutanta)\nNIM: 322024034"
    for r in cell_s2_name.paragraphs[0].runs:
        r.font.name = "Arial"
        r.font.size = Pt(9)
        
    doc.save("Invoice_BEM_FEB_Optik_Melawai.docx")
    print("Saved Invoice_BEM_FEB_Optik_Melawai.docx with dual logos.")

def create_pdf_invoice():
    import shutil
    chrome_paths = [
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
    ]
    browser_bin = next((p for p in chrome_paths if os.path.exists(p)), None)
    if browser_bin:
        html_file = os.path.abspath("invoice_optik_melawai.html")
        out_pdf = os.path.abspath("Invoice_BEM_FEB_Optik_Melawai.pdf")
        out_pdf_alt = os.path.abspath("Invoice BEM FEB UKRIDA - PT Optik Melawai.pdf")
        preview_png = os.path.abspath("invoice_preview.png")
        
        # 1. Print to PDF
        cmd_pdf = [
            browser_bin,
            "--headless=new",
            "--disable-gpu",
            "--no-pdf-header-footer",
            f"--print-to-pdf={out_pdf}",
            f"file:///{html_file}"
        ]
        try:
            subprocess.run(cmd_pdf, check=True)
            print("Saved Invoice_BEM_FEB_Optik_Melawai.pdf cleanly via browser.")
            shutil.copyfile(out_pdf, out_pdf_alt)
            print("Updated Invoice BEM FEB UKRIDA - PT Optik Melawai.pdf")
        except Exception as e:
            print("Could not generate PDF via browser:", e)

        # 2. Screenshot for preview
        cmd_shot = [
            browser_bin,
            "--headless=new",
            "--disable-gpu",
            "--window-size=900,1280",
            f"--screenshot={preview_png}",
            f"file:///{html_file}"
        ]
        try:
            subprocess.run(cmd_shot, check=True)
            print("Saved invoice_preview.png cleanly.")
        except Exception as e:
            print("Could not capture screenshot:", e)

if __name__ == "__main__":
    create_excel_invoice()
    create_word_invoice()
    create_pdf_invoice()
